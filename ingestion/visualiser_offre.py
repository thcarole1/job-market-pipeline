"""
visualiser_offre.py
-------------------
Visualise la structure d'une offre d'emploi (JSON) sous forme d'arbre
et l'exporte en fichier Word (.docx) ou PDF.

Usage :
    python visualiser_offre.py offres.json --format word
    python visualiser_offre.py offres.json --format pdf
    python visualiser_offre.py offres.json --format word --index 2
"""

import json
import argparse
from datetime import datetime
from pathlib import Path


# ─────────────────────────────────────────────
# 1. CONSTRUCTION DE L'ARBRE (texte brut)
# ─────────────────────────────────────────────

def construire_arbre(data, lignes: list, prefixe: str = "", est_dernier: bool = True):
    """
    Parcourt récursivement un dict/list JSON et construit
    une liste de lignes texte représentant l'arbre.
    """
    connecteur    = "└── " if est_dernier else "├── "
    prolongement  = "    " if est_dernier else "│   "

    if isinstance(data, dict):
        items = list(data.items())
        for i, (cle, valeur) in enumerate(items):
            dernier = (i == len(items) - 1)
            conn    = "└── " if dernier else "├── "
            prol    = "    " if dernier else "│   "

            if isinstance(valeur, dict):
                lignes.append(f"{prefixe}{conn}{cle}  [dict]")
                construire_arbre(valeur, lignes, prefixe + prol, True)

            elif isinstance(valeur, list):
                lignes.append(f"{prefixe}{conn}{cle}  [liste — {len(valeur)} élément(s)]")
                if valeur and isinstance(valeur[0], dict):
                    construire_arbre(valeur[0], lignes, prefixe + prol, True)

            else:
                type_val = type(valeur).__name__
                # Tronquer les valeurs longues pour la lisibilité
                val_str = str(valeur)
                if len(val_str) > 60:
                    val_str = val_str[:57] + "..."
                lignes.append(f"{prefixe}{conn}{cle}  ({type_val})  →  {val_str}")

    elif isinstance(data, list):
        for i, item in enumerate(data):
            dernier = (i == len(data) - 1)
            conn    = "└── " if dernier else "├── "
            prol    = "    " if dernier else "│   "
            lignes.append(f"{prefixe}{conn}[{i}]")
            if isinstance(item, (dict, list)):
                construire_arbre(item, lignes, prefixe + prol, True)


def generer_arbre_texte(offre: dict, titre: str = "Offre d'emploi") -> list:
    """
    Retourne une liste de lignes texte représentant l'arbre de l'offre.
    """
    lignes = [titre, ""]
    construire_arbre(offre, lignes)
    return lignes


# ─────────────────────────────────────────────
# 2. EXPORT WORD (.docx)
# ─────────────────────────────────────────────

def exporter_word(lignes: list, chemin_sortie: str):
    """
    Exporte l'arbre textuel dans un fichier Word (.docx).
    Utilise python-docx.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Style du document
    style = doc.styles["Normal"]
    style.font.name = "Courier New"
    style.font.size = Pt(9)

    # Titre principal
    titre_para = doc.add_heading(lignes[0], level=1)
    titre_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Horodatage
    doc.add_paragraph(
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ).runs[0].font.color.rgb = RGBColor(0x88, 0x87, 0x80)

    doc.add_paragraph("")

    # Contenu de l'arbre — chaque ligne est un paragraphe monospace
    for ligne in lignes[2:]:
        para = doc.add_paragraph()
        run  = para.add_run(ligne)
        run.font.name = "Courier New"
        run.font.size = Pt(8)

        # Coloration selon le type de nœud
        if "[dict]" in ligne:
            run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)  # bleu
        elif "[liste" in ligne:
            run.font.color.rgb = RGBColor(0xBA, 0x75, 0x17)  # amber
        elif "→" in ligne:
            run.font.color.rgb = RGBColor(0x3B, 0x6D, 0x11)  # vert
        else:
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x41)  # gris foncé

        para.paragraph_format.space_after  = Pt(0)
        para.paragraph_format.space_before = Pt(0)

    doc.save(chemin_sortie)
    print(f"Word sauvegardé : {chemin_sortie}")


# ─────────────────────────────────────────────
# 3. EXPORT PDF
# ─────────────────────────────────────────────

def exporter_pdf(lignes: list, chemin_sortie: str):
    """
    Exporte l'arbre textuel dans un fichier PDF.
    Utilise reportlab.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import mm

    doc    = SimpleDocTemplate(
        chemin_sortie,
        pagesize=A4,
        leftMargin=15*mm,
        rightMargin=15*mm,
        topMargin=15*mm,
        bottomMargin=15*mm,
    )
    styles = getSampleStyleSheet()
    story  = []

    # Style titre
    style_titre = ParagraphStyle(
        "Titre",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#1D9E75"),
        spaceAfter=4,
    )

    # Style horodatage
    style_date = ParagraphStyle(
        "Date",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.HexColor("#888780"),
        spaceAfter=10,
    )

    # Style nœud dict
    style_dict = ParagraphStyle(
        "Dict",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=7.5,
        textColor=colors.HexColor("#0C447C"),
        spaceAfter=0,
        spaceBefore=0,
        leading=11,
    )

    # Style nœud liste
    style_liste = ParagraphStyle(
        "Liste",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=7.5,
        textColor=colors.HexColor("#633806"),
        spaceAfter=0,
        spaceBefore=0,
        leading=11,
    )

    # Style feuille (valeur simple)
    style_feuille = ParagraphStyle(
        "Feuille",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=7.5,
        textColor=colors.HexColor("#27500A"),
        spaceAfter=0,
        spaceBefore=0,
        leading=11,
    )

    # Style par défaut
    style_defaut = ParagraphStyle(
        "Defaut",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=7.5,
        textColor=colors.HexColor("#444441"),
        spaceAfter=0,
        spaceBefore=0,
        leading=11,
    )

    # Titre
    story.append(Paragraph(lignes[0], style_titre))
    story.append(Paragraph(
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
        style_date
    ))

    # Contenu de l'arbre
    for ligne in lignes[2:]:
        # Échapper les caractères spéciaux HTML pour reportlab
        ligne_safe = (
            ligne
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if "[dict]" in ligne:
            style = style_dict
        elif "[liste" in ligne:
            style = style_liste
        elif "→" in ligne_safe.replace("&gt;", ">"):
            style = style_feuille
        else:
            style = style_defaut

        story.append(Paragraph(ligne_safe, style))

    doc.build(story)
    print(f"PDF sauvegardé : {chemin_sortie}")


# ─────────────────────────────────────────────
# 4. POINT D'ENTRÉE
# ─────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Visualise la structure d'une offre JSON et l'exporte en Word ou PDF."
    )
    parser.add_argument(
        "fichier_json",
        help="Chemin vers le fichier JSON contenant les offres"
    )
    parser.add_argument(
        "--format",
        choices=["word", "pdf"],
        default="word",
        help="Format de sortie : word (défaut) ou pdf"
    )
    parser.add_argument(
        "--index",
        type=int,
        default=0,
        help="Index de l'offre à visualiser dans la liste (défaut : 0)"
    )
    parser.add_argument(
        "--sortie",
        default=None,
        help="Chemin du fichier de sortie (optionnel)"
    )
    args = parser.parse_args()

    # Chargement du JSON
    chemin_json = Path(args.fichier_json)
    if not chemin_json.exists():
        print(f"Erreur : fichier introuvable — {chemin_json}")
        return

    with open(chemin_json, "r", encoding="utf-8") as f:
        donnees = json.load(f)

    # Gestion liste ou dict unique
    if isinstance(donnees, list):
        if args.index >= len(donnees):
            print(f"Erreur : index {args.index} hors limites ({len(donnees)} offres)")
            return
        offre = donnees[args.index]
    else:
        offre = donnees

    # Titre de l'arbre
    titre = offre.get("title") or offre.get("titre") or f"Offre #{args.index}"

    # Construction de l'arbre texte
    lignes = generer_arbre_texte(offre, titre=f"Structure — {titre}")

    # Affichage dans le terminal
    print("\n".join(lignes))
    print()

    # Chemin de sortie
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if args.sortie:
        chemin_sortie = args.sortie
    elif args.format == "word":
        chemin_sortie = f"structure_offre_{timestamp}.docx"
    else:
        chemin_sortie = f"structure_offre_{timestamp}.pdf"

    # Export
    if args.format == "word":
        exporter_word(lignes, chemin_sortie)
    else:
        exporter_pdf(lignes, chemin_sortie)


if __name__ == "__main__":
    main()